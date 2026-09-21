#!/usr/bin/env python3
"""Insert one 'Results: online mitigation, round 3' slide after the P3 'Results' slide
of slides/8-30_huaman_edit.pptx, in that deck's register. Backs up once to
8-30_huaman_edit.pre-p3b.bak. Text is filled from results/real/p3b_round3/minja_r3_summary.json."""
import json
import os
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIG = os.path.join(HERE, "figs_0830")
SRC = os.path.join(HERE, "8-30_huaman_edit.pptx")
BAK = os.path.join(HERE, "8-30_huaman_edit.pre-p3b.bak")
if not os.path.exists(BAK):
    shutil.copy2(SRC, BAK)
BLACK = RGBColor(0, 0, 0); BODY_FONT = "Georgia"; LEFT, CW = 0.38, 9.24

s = json.load(open(os.path.join(ROOT, "results/real/p3b_round3/minja_r3_summary.json")))
J, M = s["judgement"], s["micro"]
g1, g2 = J["g1"], J["g2"]
b1, b2 = g1["touched_paired_noop_minus_arm"], g2["touched_paired_noop_minus_arm"]

prs = Presentation(SRC)
titles = [next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "") for sl in prs.slides]
assert not any(t.startswith("Results: online mitigation") for t in titles), "slide already present"
# anchor: the P3 'Results' slide (the last slide titled exactly 'Results')
ANCHOR = max(i for i, t in enumerate(titles) if t.strip() == "Results")
BLANK = ([l for l in prs.slide_layouts if l.name.upper() == "BLANK"] or [prs.slide_layouts[6]])[0]
sl = prs.slides.add_slide(BLANK)
for ph in list(sl.placeholders):
    ph._element.getparent().remove(ph._element)
tb = sl.shapes.add_textbox(Inches(LEFT), Inches(0.18), Inches(CW), Inches(0.44))
r = tb.text_frame.paragraphs[0].add_run(); r.text = "Results: online mitigation, round 3"
r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BLACK; r.font.name = BODY_FONT
pic = sl.shapes.add_picture(os.path.join(FIG, "p3b_r3.png"), Inches(0.45), Inches(0.62), width=Inches(9.10))
y = (pic.top + pic.height) / 914400 + 0.14
tb = sl.shapes.add_textbox(Inches(LEFT), Inches(y), Inches(CW), Inches(5.45 - y))
tf = tb.text_frame; tf.word_wrap = True


def line(txt, bold=False, first=False):
    para = tf.paragraphs[0] if first else tf.add_paragraph()
    para.line_spacing = 1.3
    rr = para.add_run(); rr.text = txt; rr.font.size = Pt(12); rr.font.bold = bold; rr.font.color.rgb = BLACK; rr.font.name = BODY_FONT


def fmt(b):
    return f"{b['mean']:+.2f} [{b['ci_lo']:+.2f}, {b['ci_hi']:+.2f}], n = {b['n']}"


line("MINJA, ten seed blocks, four arms from one frozen memory (ungated, no-op, g1, g2):", True, True)
line(f"Round 2 removed only the records calibration had scored, and half the memory was never scored. Round 3 expands an "
     f"implicated record to every record written on the same question stem (g1) and, as a variant, withholds all unvetted "
     f"records in the trigger regime (g2).")
line(f"Round 3 (ten 12-round blocks): attack rate ungated {M['ungated']['asr']:.2f}, no-op {M['noop']['asr']:.2f}, g1 {M['g1']['asr']:.2f}, "
     f"g2 {M['g2']['asr']:.2f}; paired effect on touched queries against no-op g1 {fmt(b1)}, g2 {fmt(b2)}; frozen judgement "
     f"g1 {'PASS' if g1['pass'] else 'FAIL'}, g2 {'PASS' if g2['pass'] else 'FAIL'} (no block reached the evaluability floor at this base rate).")
r4p = os.path.join(ROOT, "results/real/p3b_round3/round4/minja_r4_summary.json")
if os.path.exists(r4p):
    s4 = json.load(open(r4p)); J4, M4 = s4["judgement"], s4["micro"]
    c1, c2 = J4["g1"]["touched_paired_noop_minus_arm"], J4["g2"]["touched_paired_noop_minus_arm"]
    line(f"Round 4 (fresh seeds, four three-seed blocks): ungated {M4['ungated']['asr']:.2f}, no-op {M4['noop']['asr']:.2f}, g1 {M4['g1']['asr']:.2f}, "
         f"g2 {M4['g2']['asr']:.2f}; g1 {fmt(c1)}, g2 {fmt(c2)}; evaluable blocks {J4['g1']['evaluable_blocks']} of 4; "
         f"frozen judgement g1 {'PASS' if J4['g1']['pass'] else 'FAIL'}, g2 {'PASS' if J4['g2']['pass'] else 'FAIL'}.")
sl.notes_slide.notes_text_frame.text = ("The no-op arm is the noise control: same memory, same pipeline, nothing deleted. The paired "
                                        "difference against it is the effect of the deletion policy itself. AgentPoison round 3 is blocked "
                                        "here by network access to the dev split and the encoder.")
sld = prs.slides._sldIdLst
items = list(sld); new = items[-1]; sld.remove(new); sld.insert(ANCHOR + 1, new)
prs.save(SRC)
print("saved", SRC, len(prs.slides), "slides; inserted after slide", ANCHOR + 1)
