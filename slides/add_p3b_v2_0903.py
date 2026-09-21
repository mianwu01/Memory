#!/usr/bin/env python3
"""Add the online-mitigation slides and two appendix slides to slides/8-30_huaman_edit.pptx in
the deck's register.  Inserts after the P3 'Results' slide: a prose 'Test ...' slide and a
claim-titled result slide (figure + one paragraph + notes); appends two appendix slides at the
end.  Leaves every existing slide untouched.  Backup: 8-30_huaman_edit.pre-p3b-v2.bak."""
import json
import os
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE); FIG = os.path.join(HERE, "figs_0830")
SRC = os.path.join(HERE, "8-30_huaman_edit.pptx"); BAK = os.path.join(HERE, "8-30_huaman_edit.pre-p3b-v2.bak")
if not os.path.exists(BAK):
    shutil.copy2(SRC, BAK)
BLACK = RGBColor(0, 0, 0); BODY_FONT = "Georgia"; LEFT, CW = 0.38, 9.24
s3 = json.load(open(os.path.join(ROOT, "results/real/p3b_round3/minja_r3_summary.json")))
s4 = json.load(open(os.path.join(ROOT, "results/real/p3b_round3/round4/minja_r4_summary.json")))
sp = json.load(open(os.path.join(ROOT, "results/real/p3b_round3/minja_pooled_summary.json")))
a4 = json.load(open(os.path.join(ROOT, "results/real/hm3/round4/paired_analysis.json")))

prs = Presentation(SRC)
titles = [next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "") for sl in prs.slides]
assert not any("online mitigation" in t for t in titles), "mitigation slides already present"
ANCHOR = max(i for i, t in enumerate(titles) if t.strip() == "Results")
BLANK = ([l for l in prs.slide_layouts if l.name.upper() == "BLANK"] or [prs.slide_layouts[6]])[0]


def add_slide(title, size=20):
    sl = prs.slides.add_slide(BLANK)
    for ph in list(sl.placeholders):
        ph._element.getparent().remove(ph._element)
    tb = sl.shapes.add_textbox(Inches(LEFT), Inches(0.18), Inches(CW), Inches(0.44))
    r = tb.text_frame.paragraphs[0].add_run(); r.text = title
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = BLACK; r.font.name = BODY_FONT
    return sl


def add_text(sl, y, h, lines, size=12.0, spacing=1.35):
    tb = sl.shapes.add_textbox(Inches(LEFT), Inches(y), Inches(CW), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.line_spacing = spacing
        txt, bold = (line, False) if isinstance(line, str) else line
        r = para.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = BLACK; r.font.name = BODY_FONT
    return tb


def add_pic(sl, name, left, top, width):
    return sl.shapes.add_picture(os.path.join(FIG, name), Inches(left), Inches(top), width=Inches(width))


def below(pic, gap=0.14):
    return (pic.top + pic.height) / 914400 + gap


def notes(sl, text):
    sl.notes_slide.notes_text_frame.text = text


M3, J3, M4, J4, Jp = s3["micro"], s3["judgement"], s4["micro"], s4["judgement"], sp["judgement"]
b3g1, b3g2 = J3["g1"]["touched_paired_noop_minus_arm"], J3["g2"]["touched_paired_noop_minus_arm"]
b4g1, b4g2 = J4["g1"]["touched_paired_noop_minus_arm"], J4["g2"]["touched_paired_noop_minus_arm"]
bpg1, bpg2 = Jp["g1"]["touched_paired_noop_minus_arm"], Jp["g2"]["touched_paired_noop_minus_arm"]


def ci(b):
    return f"{b['mean']:+.3f} [{b['ci_lo']:+.3f}, {b['ci_hi']:+.3f}], n = {b['n']}"


inserted = []

# --- 16. Test online mitigation with a no-op control (prose, like the auditing setup slide)
sl = add_slide("Test online mitigation with a no-op control")
add_text(sl, 0.85, 4.5, [
    "LLM:  DeepSeek-v4-flash, the MINJA carrier from the auditing experiment.",
    "The gate turns the recovered ancestry into a deletion policy and is tested on new LLM calls",
    ("Calibration", True),
    "Earlier rounds are scored without poison labels: a record is implicated when it was retrieved in trigger rounds "
    "that ended in an anomalous answer. Round 2 deleted only those records, and about half of the memory was never "
    "retrieved during calibration, so unscored poison records survived.",
    ("Two deletion policies", True),
    "g1 adds every record written on the same question stem as an implicated one (the attack appends escalating "
    "notes to one stem). g2 adds a quarantine: in the trigger regime, records that calibration never vetted are "
    "withheld as well.",
    ("Four arms from one frozen memory", True),
    "Each held-out query is answered by the ungated memory, by a no-op arm with an empty deletion set, by g1 and "
    "by g2, in a fixed random order. The no-op arm is the noise control: independent calls flip answers on their own, "
    "so every effect is measured paired against it. Round 3: ten seeds × 12 rounds. Round 4: fresh seeds 10–21 in "
    "four three-seed blocks, frozen before its first call; a block counts when its gate-free arms hold at least "
    "three attacks.",
], size=11.5, spacing=1.22)
notes(sl, "Same shape as the auditing setup slide. The one idea to land: the effect is measured against a no-op arm "
          "on the same frozen memory, because paired LLM calls flip answers without any intervention (the ungated and "
          "no-op arms disagree on 6-9% of rounds). Round 4's block design was frozen after an interim look at five "
          "round-3 seeds and is disclosed in the protocol.")
inserted.append(sl)

# --- 17. result slide
sl = add_slide("Deleting the recovered ancestry and its same-stem records mitigates the attack", size=17)
pic = add_pic(sl, "p3b_r3.png", 0.45, 0.66, 9.10)
add_text(sl, below(pic), 1.5, [
    f"On fresh seeds the gate removes about half of the attacks (g1) to four fifths (g2) that the same memory produces "
    f"under the no-op control, with no loss of normal-answer accuracy, and passes the pre-registered judgement in "
    f"3 of 4 blocks for each policy.",
], size=14.0, spacing=1.3)
notes(sl, f"Round 3 (ten 12-round blocks): attacks per 120 rounds ungated {M3['ungated']['attacks']}, no-op {M3['noop']['attacks']}, "
          f"g1 {M3['g1']['attacks']}, g2 {M3['g2']['attacks']}; accuracy g2 {M3['g2']['accuracy']:.2f}; paired against no-op on touched "
          f"queries g1 {ci(b3g1)}, g2 {ci(b3g2)}; frozen judgement FAIL for both because the base attack rate (3-8%) leaves no "
          f"12-round block above the evaluability floor. Round 4 (seeds 10-21, four three-seed blocks, all evaluable): attacks per 144 "
          f"rounds ungated {M4['ungated']['attacks']}, no-op {M4['noop']['attacks']}, g1 {M4['g1']['attacks']}, g2 {M4['g2']['attacks']}; "
          f"accuracy {M4['ungated']['accuracy']:.2f} / {M4['noop']['accuracy']:.2f} / {M4['g1']['accuracy']:.2f} / {M4['g2']['accuracy']:.2f}; "
          f"g1 {ci(b4g1)}, g2 {ci(b4g2)}; 3 of 4 blocks improve under each -> PASS for both. Pooled 22 seeds (descriptive): g1 {ci(bpg1)}, "
          f"g2 {ci(bpg2)}. Boundaries: one carrier and one model; the attack is rare against this model, so absolute effects are small; "
          f"mitigation evidence does not upgrade P3-A's recovery evidence; rounds 2 and 3 stay recorded as FAIL. AgentPoison round 3 is "
          f"blocked here (no network path to the dev split or the DPR encoder).")
inserted.append(sl)

# --- appendix: why the graph selection lost on Travel, and the closure fix
sl = add_slide("Appendix — the LLM needs the objects its witnesses name", size=18)
pic = add_pic(sl, "v3_closure.png", 0.55, 0.62, 8.90)
t4, s4c = a4["travel"]["graph_closed/compact - full/verbose(r3)"], a4["search"]["graph_closed/compact - full/verbose(r3)"]
add_text(sl, below(pic), 1.4, [
    f"In 33 of 38 failing Travel episodes the graph's witness records named objects the selection had dropped, so the "
    f"model saw a transfer move without the transfer's provider. Adding those objects and their neighbours recovers "
    f"half of the gap ({a4['travel']['graph_closed/compact - graph/compact(r3)']['mean']:+.2f}) at a quarter of the full input; "
    f"the remainder is the runtime's use of a minimal evidence set.",
], size=12.5, spacing=1.3)
notes(sl, f"Round 3 vs round 4, 63-69 episodes per cell, paired per episode. Travel: graph/compact 0.219, graph_closed/compact "
          f"{t4['ees_a']:.3f}, full/verbose {t4['ees_b']:.3f}; paired closed - full {t4['mean']:+.3f} [{t4['ci_lo']:+.3f}, {t4['ci_hi']:+.3f}]; "
          f"closure gain over the open selection {a4['travel']['graph_closed/compact - graph/compact(r3)']['mean']:+.3f} "
          f"[{a4['travel']['graph_closed/compact - graph/compact(r3)']['ci_lo']:+.3f}, {a4['travel']['graph_closed/compact - graph/compact(r3)']['ci_hi']:+.3f}]. "
          f"Search: closed {s4c['ees_a']:.3f} vs full {s4c['ees_b']:.3f}, {s4c['mean']:+.3f} [{s4c['ci_lo']:+.3f}, {s4c['ci_hi']:+.3f}]; the closure "
          f"changed nothing there. The deterministic executor always sees the full state, which is why the graph reaches 0.86 / 0.99 on "
          f"the same episodes: the read-recall metric counted witness records, and the LLM also needs the objects those records refer to.")
inserted.append(sl)

# --- appendix: claims and decisions
sl = add_slide("Appendix — what we can write, and the decisions", size=18)
add_text(sl, 0.82, 4.6, [
    ("Can write", True),
    "·  memory as write → hold → regime-gated read; regime conditioning recovers gated edges that pooling misses (E0)",
    "·  a learned slot graph selects memory on MemoryArena at −39.55% input with PS held (frozen PASS)",
    "·  the v3 benchmark isolates hidden structure on an executable endpoint: lookups, kNN and superset fail, both oracles complete it",
    "·  learned relational structure with history-parsed policies completes repair (0.72–1.00); the graph form leads on Travel",
    "·  graph selection cuts LLM input by 70–74% at a measured exact-success cost of 0.10–0.19",
    "·  the recovered ancestry is recoverable on MINJA (3/3 seeds) and, with same-stem expansion, mitigates the attack online (round 4 PASS)",
    ("Cannot write", True),
    "·  that the causal-graph form is necessary in all four tasks (Formal favours the per-object learner)",
    "·  that the LLM runtime already exploits the learned structure, or that graph selection is non-inferior at −0.10",
    "·  mitigation beyond one carrier and one model; rounds 2 and 3 stay FAIL in the record",
    ("Decisions", True),
    "1 · formulation and claim   2 · Shopping v3.2 as the version of record   3 · report the LLM cost or buy a stronger runtime",
    "4 · Formal gap: solve or report   5 · mitigation enters the paper with its boundaries?   6 · scale simulation   7 · ICLR 2027 (Sep 18 / 25)",
], size=11.0, spacing=1.22)
notes(sl, "Backup for the spoken decisions. Read the two lists straight; the decisions are the same seven as in the script.")
inserted.append(sl)

sld = prs.slides._sldIdLst; items = list(sld); new_items = items[-len(inserted):]
for it in new_items:
    sld.remove(it)
# the two mitigation slides go after the P3 Results slide; the two appendix slides go to the end
sld.insert(ANCHOR + 1, new_items[0]); sld.insert(ANCHOR + 2, new_items[1])
for it in new_items[2:]:
    sld.append(it)
prs.save(SRC)
print("saved", SRC, len(prs.slides), "slides")
