#!/usr/bin/env python3
"""Refresh the existing 'Results: online mitigation, round 3' slide in place (picture and body
text) from the round-3 and round-4 summaries, leaving every other slide untouched."""
import json
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE); FIG = os.path.join(HERE, "figs_0830")
SRC = os.path.join(HERE, "8-30_huaman_edit.pptx")
BLACK = RGBColor(0, 0, 0); BODY_FONT = "Georgia"; LEFT, CW = 0.38, 9.24
s = json.load(open(os.path.join(ROOT, "results/real/p3b_round3/minja_r3_summary.json"))); J, M = s["judgement"], s["micro"]
r4p = os.path.join(ROOT, "results/real/p3b_round3/round4/minja_r4_summary.json")
s4 = json.load(open(r4p)) if os.path.exists(r4p) else None
prs = Presentation(SRC)
slide = next(sl for sl in prs.slides if any(sh.has_text_frame and "online mitigation" in sh.text_frame.text for sh in sl.shapes))
title = next(sh for sh in slide.shapes if sh.has_text_frame and "online mitigation" in sh.text_frame.text)
for sh in list(slide.shapes):
    if sh._element is not title._element:
        sh._element.getparent().remove(sh._element)
if s4:
    r = title.text_frame.paragraphs[0].runs[0]; r.text = "Results: online mitigation, rounds 3 and 4"
pic = slide.shapes.add_picture(os.path.join(FIG, "p3b_r3.png"), Inches(0.45), Inches(0.62), width=Inches(9.10))
y = (pic.top + pic.height) / 914400 + 0.14
tb = slide.shapes.add_textbox(Inches(LEFT), Inches(y), Inches(CW), Inches(5.45 - y)); tf = tb.text_frame; tf.word_wrap = True


def line(txt, bold=False, first=False):
    para = tf.paragraphs[0] if first else tf.add_paragraph(); para.line_spacing = 1.25
    rr = para.add_run(); rr.text = txt; rr.font.size = Pt(11.5); rr.font.bold = bold; rr.font.color.rgb = BLACK; rr.font.name = BODY_FONT


def fmt(b):
    return f"{b['mean']:+.2f} [{b['ci_lo']:+.2f}, {b['ci_hi']:+.2f}], n = {b['n']}"


b1, b2 = J["g1"]["touched_paired_noop_minus_arm"], J["g2"]["touched_paired_noop_minus_arm"]
line("MINJA, four arms from one frozen memory: ungated, no-op (noise control), g1 = implicated records plus every record "
     "written on the same question stem, g2 = g1 plus quarantine of unvetted records in the trigger regime.", True, True)
line(f"Round 3 (ten 12-round blocks): attack rate ungated {M['ungated']['asr']:.2f}, no-op {M['noop']['asr']:.2f}, g1 {M['g1']['asr']:.2f}, "
     f"g2 {M['g2']['asr']:.2f}; paired effect against no-op on touched queries g1 {fmt(b1)}, g2 {fmt(b2)}; frozen judgement "
     f"g1 {'PASS' if J['g1']['pass'] else 'FAIL'}, g2 {'PASS' if J['g2']['pass'] else 'FAIL'}: no block reached the evaluability floor at this base rate.")
if s4:
    J4, M4 = s4["judgement"], s4["micro"]; c1, c2 = J4["g1"]["touched_paired_noop_minus_arm"], J4["g2"]["touched_paired_noop_minus_arm"]
    line(f"Round 4 (fresh seeds, four three-seed blocks): ungated {M4['ungated']['asr']:.2f}, no-op {M4['noop']['asr']:.2f}, g1 {M4['g1']['asr']:.2f}, "
         f"g2 {M4['g2']['asr']:.2f}; g1 {fmt(c1)}, g2 {fmt(c2)}; evaluable blocks {J4['g1']['evaluable_blocks']} of 4; frozen judgement "
         f"g1 {'PASS' if J4['g1']['pass'] else 'FAIL'}, g2 {'PASS' if J4['g2']['pass'] else 'FAIL'}.")
prs.save(SRC); print("P3-B slide refreshed in place")
