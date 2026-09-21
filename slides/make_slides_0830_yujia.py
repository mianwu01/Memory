#!/usr/bin/env python3
"""Yujia sync deck (2026-08-30, v4): causal-memory round-1 results.

Text register follows the human-revised real meeting deck
(TTS/slides/"8-12 (1).pptx"): short plain titles, title + figure + one short
paragraph per slide, no footers on main slides, caveats in one line, arrows
and short declaratives. Appendix slides stay dense (backup only).

All numbers from on-disk sources:
  docs/HANDOFF.md · results/regime_grace_e0.json
  results/diagnostics/e0_blind_power_sweep.json   (blind: 0/2 @3k -> 2/2 @60k)
  results/real/p2_compact_v3/round_summary.json   (episode-mean + pooled deltas)
  results/real/minja_replication_summary.json · minja_online_gate_summary.json
  results/real/p3_agentpoison_round2/summary.json · causal_gate_v2.json
  results/real/p3b_power_audit.json
Estimators verified from code: GRACE-per-regime ran P2's slot graph
(travel_grace_discovery.py); per-regime ridge + BH-FDR ran E0 and MINJA
(regime_grace.py, minja_causal_analysis.py).
ICLR 2027 CFP (iclr.cc): abstracts Sep 18, papers Sep 25, 2026 AOE.
"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figs_0830")
os.makedirs(FIG, exist_ok=True)

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x55, 0x55, 0x55)
BODY_FONT = "Georgia"

plt.rcParams.update({
    "figure.facecolor": "white", "axes.facecolor": "white",
    "axes.edgecolor": "black", "axes.labelcolor": "black",
    "xtick.color": "black", "ytick.color": "black",
    "text.color": "black", "font.size": 10,
    "font.family": "serif", "font.serif": ["Georgia", "DejaVu Serif"],
    "axes.spines.top": False, "axes.spines.right": False,
})

BLUE = "#2a78d6"
LIGHTBLUE = (0.88, 0.92, 0.99)


# =====================================================================
# FIGURES
# =====================================================================

def fig_object():
    """The primitive chain + one learned graph, two uses."""
    f, ax = plt.subplots(figsize=(9.2, 3.1))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")

    def box(x, y, w, h, txt, fc="white", lw=1.4, fs=10):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.010",
                                    fc=fc, ec="black", lw=lw))
        ax.text(x + w / 2, y + h / 2, txt, ha="center", va="center", fontsize=fs)

    def arrow(x1, y1, x2, y2, lw=1.4, color="black", ls="-", rad=0.0):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
                                     mutation_scale=14, lw=lw, color=color,
                                     ls=ls, connectionstyle=f"arc3,rad={rad}"))

    cy, ch = 0.72, 0.22
    box(0.02, cy, 0.135, ch, "write\n$X_{t-k}$", fc=LIGHTBLUE, fs=10)
    arrow(0.158, cy + ch / 2, 0.198, cy + ch / 2)
    box(0.20, cy, 0.105, ch, "$M_{t-k}$")
    arrow(0.308, cy + ch / 2, 0.348, cy + ch / 2)
    box(0.35, cy, 0.105, ch, "$M_{t-k+1}$")
    arrow(0.458, cy + ch / 2, 0.498, cy + ch / 2)
    ax.text(0.525, cy + ch / 2, r"$\cdots$", ha="center", va="center", fontsize=13)
    arrow(0.552, cy + ch / 2, 0.592, cy + ch / 2)
    box(0.594, cy, 0.095, ch, "$M_t$")
    arrow(0.692, cy + ch / 2, 0.792, cy + ch / 2, lw=2.0, color=BLUE)
    ax.text(0.742, cy + ch + 0.035, "$U_t=\\mathrm{read}$", ha="center",
            fontsize=10.5, color=BLUE)
    ax.text(0.742, cy - 0.10, "gate: edge active only\nin the read regime",
            ha="center", va="top", fontsize=8.2, style="italic", color="0.3")
    box(0.795, cy, 0.095, ch, "$Y_t$", fc=LIGHTBLUE, fs=11)
    ax.text(0.295, cy - 0.10, "hold: value persists,\nedge inactive",
            ha="center", va="top", fontsize=8.2, style="italic", color="0.3")

    box(0.40, 0.30, 0.20, 0.155, "learned graph  $\\hat G$\n(edges + gates)",
        lw=1.7, fs=9.6)
    arrow(0.50, 0.70, 0.50, 0.465, lw=1.2, ls=(0, (4, 2)))
    ax.text(0.515, 0.575, "discovery (E0)", fontsize=8.6, style="italic")

    box(0.03, 0.02, 0.34, 0.20,
        "forward — memory selection (P2)\n$\\hat G\\rightarrow F_t$: which cells "
        "the current query needs",
        fs=9.2)
    box(0.63, 0.02, 0.34, 0.20,
        "backward — memory auditing (P3-A)\n$Y^{bad}\\rightarrow\\hat G\\rightarrow$ "
        "ancestry: the write it came from",
        fs=9.2)
    arrow(0.42, 0.31, 0.235, 0.235, lw=1.4, rad=0.12)
    arrow(0.58, 0.31, 0.765, 0.235, lw=1.4, rad=-0.12)

    f.tight_layout(pad=0.2)
    p = os.path.join(FIG, "object.png")
    f.savefig(p, dpi=200, bbox_inches="tight"); plt.close(f); return p


def fig_regime():
    """E0: write-hold-read timeline + read-edge recovery bars."""
    f, (ax1, ax2) = plt.subplots(1, 2, figsize=(9.3, 2.75),
                                 gridspec_kw={"width_ratios": [1.25, 1]})
    ax1.set_xlim(-0.7, 14.7); ax1.set_ylim(0, 1); ax1.axis("off")
    ax1.plot([0, 14], [0.62, 0.62], color="black", lw=1.2)
    for t in range(15):
        ax1.plot([t, t], [0.595, 0.645], color="black", lw=0.8)
    ax1.plot([0], [0.62], "s", color=BLUE, ms=9)
    ax1.text(0, 0.76, "write\n$t-\\Delta$", ha="center", fontsize=9)
    ax1.plot([10], [0.62], "o", color=BLUE, ms=9)
    ax1.text(10, 0.76, "read\n$U_t=1$", ha="center", fontsize=9)
    ax1.annotate("", xy=(9.6, 0.50), xytext=(0.4, 0.50),
                 arrowprops=dict(arrowstyle="<->", lw=1.0, color="0.35"))
    ax1.text(5, 0.40, "hold, $\\Delta=10$ steps — value persists, edge inactive",
             ha="center", fontsize=8.4, style="italic", color="0.3")
    ax1.plot([0, 9.8], [0.12, 0.12], color="0.45", lw=1.6)
    ax1.plot([9.8, 9.8], [0.12, 0.24], color=BLUE, lw=1.6)
    ax1.plot([9.8, 10.2], [0.24, 0.24], color=BLUE, lw=2.4)
    ax1.plot([10.2, 10.2], [0.24, 0.12], color=BLUE, lw=1.6)
    ax1.plot([10.2, 14], [0.12, 0.12], color="0.45", lw=1.6)
    ax1.text(-0.6, 0.17, "$u_t$", fontsize=10, ha="right")
    ax1.text(12.4, 0.245, "read regime: 150 / 3,000\nsamples (5%)",
             fontsize=8.2, color="0.3", va="bottom", ha="center")
    ax1.set_title("pooled over all steps, the read edge's\nsignal falls below "
                  "the test threshold", fontsize=9.5)
    labs = ["Regime-GRACE\n(gate in the\nmechanism)", "blind PCMCI+\n(regime ignored)",
            "additive $u$\n(regime as a node)"]
    vals = [2, 0, 0]
    xs = range(3)
    ax2.bar(xs, vals, color=[BLUE, "white", "white"], edgecolor="black", width=0.55)
    for x, v in zip(xs, vals):
        ax2.text(x, v + 0.07, f"{v}/2", ha="center", fontsize=11,
                 fontweight="bold" if x == 0 else "normal")
    ax2.set_xticks(list(xs)); ax2.set_xticklabels(labs, fontsize=8.4)
    ax2.set_ylabel("gated read edges recovered")
    ax2.set_ylim(0, 2.5); ax2.set_yticks([0, 1, 2])
    ax2.set_title("3,000 steps · same at every\n$\\sigma$ · gate localized 2/2",
                  fontsize=9.5)
    f.tight_layout()
    p = os.path.join(FIG, "regime.png")
    f.savefig(p, dpi=200, bbox_inches="tight"); plt.close(f); return p


def fig_p2():
    """P2 frozen IDs 111-120: PS, input tokens, frozen criterion."""
    f = plt.figure(figsize=(9.3, 2.95))
    gs = f.add_gridspec(2, 3, width_ratios=[1.12, 1.12, 1.0], hspace=0.9,
                        wspace=0.45)
    axA = f.add_subplot(gs[:, 0]); axB = f.add_subplot(gs[:, 1])
    axC1 = f.add_subplot(gs[0, 2]); axC2 = f.add_subplot(gs[1, 2])

    labs = ["query-\nancestry\ngraph", "noG", "BM25", "long\ncontext"]
    cols = [BLUE, "white", "white", "white"]
    xs = range(4)
    ps = [93.94, 95.45, 72.73, 92.42]
    axA.bar(xs, ps, color=cols, edgecolor="black", width=0.6)
    for x, v in zip(xs, ps):
        axA.text(x, v + 2.2, f"{v:.2f}", ha="center", fontsize=9.5,
                 fontweight="bold" if x == 0 else "normal")
    axA.set_xticks(list(xs)); axA.set_xticklabels(labs, fontsize=8.4)
    axA.set_ylabel("plan score PS (%), all persons pooled")
    axA.set_ylim(0, 112)
    axA.set_title("quality is preserved", fontsize=9.5)
    inp = [434, 718, 1458, 2065]
    axB.bar(xs, inp, color=cols, edgecolor="black", width=0.6)
    for x, v in zip(xs, inp):
        axB.text(x, v + 45, f"{v}k", ha="center", fontsize=9.5,
                 fontweight="bold" if x == 0 else "normal")
    axB.set_xticks(list(xs)); axB.set_xticklabels(labs, fontsize=8.4)
    axB.set_ylabel("API input tokens (thousands)")
    axB.set_ylim(0, 2320)
    axB.set_title("while context shrinks", fontsize=9.5)
    axC1.errorbar([-1.43], [0], xerr=[[abs(-4.29 - (-1.43))], [abs(0.00 - (-1.43))]],
                  fmt="o", color=BLUE, ms=7, lw=1.6, capsize=3)
    axC1.axvline(-5.0, ls="--", color="black", lw=1.2)
    axC1.text(-4.92, 0.62, "frozen floor $-5$", fontsize=8.2)
    axC1.text(-1.43, -0.72, "$-1.43$", ha="center", fontsize=9.5, fontweight="bold")
    axC1.set_xlim(-6.6, 0.9); axC1.set_ylim(-1, 1)
    axC1.set_yticks([])
    axC1.set_xlabel("episode-mean $\\Delta$PS vs noG\n(points, 95% CI)", fontsize=8.0)
    axC2.barh([0], [39.55], color=BLUE, edgecolor="black", height=0.5)
    axC2.axvline(30.0, ls="--", color="black", lw=1.2)
    axC2.text(29.2, 0.58, "required $\\geq 30\\%$", fontsize=8.2, ha="right")
    axC2.text(14.5, 0, "39.55%", ha="center", va="center", fontsize=9.5,
              fontweight="bold", color="white")
    axC2.set_xlim(0, 50); axC2.set_ylim(-0.9, 1.1)
    axC2.set_yticks([])
    axC2.set_xlabel("input reduction vs noG (%)", fontsize=8.2)
    axC2.text(49.5, 0.62, "PASS", fontsize=15, fontweight="bold", ha="right",
              color=BLUE)
    axC1.set_title("frozen conjunctive criterion", fontsize=9.5)
    p = os.path.join(FIG, "p2.png")
    f.savefig(p, dpi=200, bbox_inches="tight"); plt.close(f); return p


def fig_p3():
    """P3: auditing (supported) vs online mitigation (FAIL)."""
    f, (ax1, ax2) = plt.subplots(1, 2, figsize=(9.3, 2.85),
                                 gridspec_kw={"width_ratios": [1.15, 1]})
    labs = ["Regime-\nGRACE", "blind\nPCMCI", "additive\n$u$-node", "regime-\nsubsampl.",
            "driver\ndirect", "+ frozen\nexpansion"]
    vals = [1.0, 0.0, 2 / 3, 1 / 3, 0.5, 1.0]
    tags = ["3/3", "0/3", "2/3", "1/3", "1/2", "2/2"]
    xs = [0, 1, 2, 3, 4.7, 5.7]
    cols = [BLUE, "white", "white", "white", "white", BLUE]
    ax1.bar(xs, vals, color=cols, edgecolor="black", width=0.6)
    for x, v, t in zip(xs, vals, tags):
        ax1.text(x, v + 0.035, t, ha="center", fontsize=9.5,
                 fontweight="bold" if x in (0, 5.7) else "normal")
    ax1.set_xticks(xs); ax1.set_xticklabels(labs, fontsize=7.8)
    ax1.set_ylabel("recovered fraction")
    ax1.set_ylim(0, 1.46); ax1.set_yticks([0, 0.5, 1.0])
    ax1.text(1.5, 1.31, "MINJA: seeds recovering the hidden\nread edge "
             "$poison\\_retr\\rightarrow anomalous$", ha="center", va="center",
             fontsize=8.4)
    ax1.text(5.2, 1.31, "AgentPoison: poison\nrecords, label-free", ha="center",
             va="center", fontsize=8.4)
    ax1.axvline(3.85, color="0.75", lw=0.9, ls=":")
    ax1.set_title("P3-A auditing — supported", fontsize=10, pad=10)
    labs2 = ["ungated", "gated", "ungated", "no-op", "gated"]
    vals2 = [6 / 36 * 100, 5 / 36 * 100, 3 / 72 * 100, 4 / 72 * 100, 0.0]
    tags2 = ["6/36", "5/36", "3/72", "4/72", "0/72"]
    xs2 = [0, 1, 2.7, 3.7, 4.7]
    cols2 = ["white", BLUE, "white", "0.8", BLUE]
    ax2.bar(xs2, vals2, color=cols2, edgecolor="black", width=0.62)
    for x, v, t in zip(xs2, vals2, tags2):
        ax2.text(x, v + 0.55, t, ha="center", fontsize=9.5)
    ax2.set_xticks(xs2); ax2.set_xticklabels(labs2, fontsize=8.4)
    ax2.set_ylabel("held-out attack success (%)")
    ax2.set_ylim(0, 22.5)
    ax2.text(0.5, 20.6, "MINJA", ha="center", fontsize=9)
    ax2.text(3.7, 20.6, "AgentPoison-StrategyQA", ha="center", fontsize=9)
    ax2.axvline(1.85, color="0.75", lw=0.9, ls=":")
    ax2.text(3.7, 12.6, "ASR drops on both, but only\n1/3 seed blocks improve;\n"
             "frozen $\\geq 2/3$ rule: FAIL",
             ha="center", fontsize=8.6, style="italic")
    ax2.set_title("P3-B online mitigation — FAIL", fontsize=10, pad=10)
    f.tight_layout()
    p = os.path.join(FIG, "p3.png")
    f.savefig(p, dpi=200, bbox_inches="tight"); plt.close(f); return p


def fig_gate():
    """Appendix: replay-only gate coverage."""
    f, (ax1, ax2) = plt.subplots(1, 2, figsize=(9.0, 2.1), sharey=True)
    labs = ["ancestry + regime", "ancestry, no regime", "CMIS + regime"]
    ys = [2, 1, 0]
    prev = [95.3, 97.2, 36.4]
    coll = [8.8, 40.4, 2.9]
    cols = [BLUE, "white", "white"]
    ax1.barh(ys, prev, color=cols, edgecolor="black", height=0.55)
    for y, v in zip(ys, prev):
        ax1.text(v - 2, y, f"{v}%", va="center", ha="right", fontsize=9.5,
                 color="white" if v > 90 and y == 2 else "black")
    ax1.set_yticks(ys); ax1.set_yticklabels(labs, fontsize=9)
    ax1.set_xlim(0, 105)
    ax1.set_xlabel("anomalous rounds covered (102-104 of 107)")
    ax1.set_title("prevention coverage (higher is better)", fontsize=9.5)
    ax2.barh(ys, coll, color=cols, edgecolor="black", height=0.55)
    for y, v in zip(ys, coll):
        ax2.text(v + 1.2, y, f"{v}%", va="center", fontsize=9.5)
    ax2.set_xlim(0, 50)
    ax2.set_xlabel("benign rounds touched (of 136)")
    ax2.set_title("collateral (lower is better)", fontsize=9.5)
    f.tight_layout()
    p = os.path.join(FIG, "gate.png")
    f.savefig(p, dpi=200, bbox_inches="tight"); plt.close(f); return p


# =====================================================================
# DECK
# =====================================================================
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)
BLANK = ([l for l in prs.slide_layouts if l.name.upper() == "BLANK"] or
         [prs.slide_layouts[6]])[0]
LEFT, CW = 0.38, 9.24


def add_slide(title=None, title_size=20, tag=None):
    s = prs.slides.add_slide(BLANK)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    if title:
        tb = s.shapes.add_textbox(Inches(LEFT), Inches(0.18), Inches(CW), Inches(0.55))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = title; r.font.size = Pt(title_size); r.font.bold = True
        r.font.color.rgb = BLACK; r.font.name = BODY_FONT
    if tag:
        tb = s.shapes.add_textbox(Inches(6.4), Inches(0.02), Inches(3.35), Inches(0.3))
        p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = tag
        r.font.size = Pt(10); r.font.italic = True
        r.font.color.rgb = GRAY; r.font.name = BODY_FONT
    return s


def add_text(slide, x, y, w, h, lines, size=11, leading=1.12, color=BLACK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.line_spacing = leading
        if isinstance(line, tuple) and len(line) == 4:
            txt, sz, bold, col = line
        elif isinstance(line, tuple):
            txt, sz, bold = line; col = color
        else:
            txt, sz, bold, col = line, size, False, color
        r = para.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold
        r.font.color.rgb = col; r.font.name = BODY_FONT
    return tb


def footer2(slide, text):
    """Appendix-only source footer."""
    add_text(slide, LEFT, 5.14, CW, 0.42, [(text, 9, False)], color=GRAY)


def pic(slide, path, x, y, w):
    slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


F_OBJECT = fig_object()
F_REGIME = fig_regime()
F_P2 = fig_p2()
F_P3 = fig_p3()
F_GATE = fig_gate()

# ---------------- S0: title ----------------
s = add_slide()
add_text(s, LEFT, 0.75, CW, 1.2, [
    ("Causal Memory", 32, True),
    ("Round 1 results · 2026-08-30", 14, False, GRAY)], leading=1.3)
add_text(s, LEFT, 2.05, CW, 0.9, [
    ("Memory = write → hold → regime-gated read  (Layer 1, observed)", 15, False),
    ("One learned graph, two uses: forward memory selection · backward auditing",
     15, False)], leading=1.5)
add_text(s, LEFT, 3.25, CW, 1.4, [
    ("E0 simulation — gated-edge recovery: supported", 12.5, False),
    ("P2 MemoryArena — compact memory: PASS (frozen criterion)", 12.5, False),
    ("P3-A auditing — supported, with boundaries", 12.5, False),
    ("P3-B online mitigation — FAIL (kept)", 12.5, False)], leading=1.55)
add_text(s, LEFT, 4.72, CW, 0.5, [
    ("Today: v0.2 formulation · samples definition · scaling simulation · venue",
     13, True)], leading=1.3)
notes(s, "Round 1 is complete and frozen: two supported claims, one PASS, one "
         "kept negative. The meeting decides the four open items; the last slide "
         "turns them into six questions.")

# ---------------- S1: the object ----------------
s = add_slide("Memory = write → hold → regime-gated read", 20, tag="the object")
pic(s, F_OBJECT, 0.55, 0.85, 8.9)
add_text(s, LEFT, 4.75, CW, 0.6, [
    ("The same learned graph selects memory for the current query (forward) and "
     "traces a bad output to the write that caused it (backward).", 12.5, False)],
    leading=1.3)
notes(s, "The one-picture story. E0 shows the graph is recoverable, P2 uses it "
         "forward, P3-A uses it backward. This is the 8/21 'one structure, two "
         "payoffs' framing.")

# ---------------- S2: 8/15 questions ----------------
s = add_slide("Your 8/15 questions → v0.2", 20, tag="formulation")
add_text(s, LEFT, 0.85, 4.2, 3.5, [
    ("Your questions:", 12, True),
    ("·  how many variables can GRACE handle?", 11.5, False),
    ("·  what is a variable? what granularity?", 11.5, False),
    ("·  what is a sample? how many per variable?", 11.5, False),
    ("·  what does the graph learn / how does it act on memory?", 11.5, False),
    ("·  trustworthiness or downstream benefit?", 11.5, False)], leading=1.55)
add_text(s, 4.85, 0.85, 4.8, 3.7, [
    ("v0.2:", 12, True),
    ("·  t = interaction step; Xₜ = semantic slots; unwritten slots persist "
     "(Xₜ = Xₜ₋₁)", 11, False),
    ("·  Uₜ = read/write regime; discovery input (X, U) → graph + per-edge gates",
     11, False),
    ("·  sample = one episode (independent trial); effective samples counted per "
     "edge × lag × regime", 11, False),
    ("·  forward: graph → what memory to keep; backward: graph → ancestry of a "
     "bad output", 11, False),
    ("·  downstream score primary; trustworthiness = second use of the same graph",
     11, False)], leading=1.5)
add_text(s, LEFT, 4.72, CW, 0.5, [
    ("Scale and the written-out SCM: next two slides.", 12.5, True)], leading=1.3)
notes(s, "Variable definition uses one of the three options she listed (vector "
         "slots allowed); she left the choice open, so this is for sign-off — "
         "decisions 1 and 2.")

# ---------------- S3: SCM ----------------
s = add_slide("The formulation (Layer 1)", 20, tag="SCM & method")
add_text(s, LEFT, 0.85, CW, 3.9, [
    ("Xₜ = (Xₜ¹ … Xₜᵈ): semantic slots, observed  ·  Uₜ ∈ {write, hold, read}  ·  "
     "Yₜ: output", 12, False),
    ("", 6, False),
    ("xⱼ(t)  =  Σᵢ Σₗ₌₀…ᴷ [ aᵢⱼₗ · gᵢⱼₗ(uₜ) ] · xᵢ(t−l)  +  εⱼ(t)", 16, True),
    ("l = 0: within-step · l ≥ 1: cross-step · hold: unwritten slot copies itself",
     11.5, False),
    ("", 6, False),
    ("The gate multiplies the coefficient. A linear-additive uₜ-node only shifts "
     "the mean, so it misses gated edges. (A CD-NOD-style nonparametric surrogate "
     "could detect them; our round-1 baselines are linear.)", 11.5, False),
    ("", 6, False),
    ("Estimator — fit each regime separately; edge = significant in any regime; "
     "gated = coefficients differ ≥ 4×. Ran as GRACE-per-regime (P2) and "
     "per-regime ridge + FDR (E0, MINJA).", 11.5, False),
    ("", 6, False),
    ("Layer 2 = the same object over latent state (“what the agent is thinking”). "
     "Round 1 is entirely Layer 1.", 11.5, False)], leading=1.4)
notes(s, "The functions she asked for on 7/29: within-step (l=0) and cross-step "
         "(l>=1) in one equation. If she raises CD-NOD: agree — our additive arms "
         "are linear, the nonparametric surrogate is the fair comparison and is "
         "in the round-2 simulation.")

# ---------------- S4: scale ----------------
s = add_slide("Scale", 20, tag="your first 8/15 question")
add_text(s, LEFT, 0.90, CW, 1.6, [
    ("Round 1 is small:  E0 8 variables · P2 7 slots · MINJA 6 event channels — "
     "all far under GRACE's ~100-variable ceiling; grouping never needed.",
     12.5, False),
    ("GRACE ran P2's slot graph; per-regime ridge ran E0 and MINJA. The ceiling "
     "belongs to the GRACE path.", 12.5, False)], leading=1.5)
add_text(s, LEFT, 2.55, CW, 0.7, [
    ("Missing: the simulation you called mandatory — SCM/MLP data, known ground "
     "truth, SHD/F1, up to ~5,000 dims.", 13, True)], leading=1.35)
add_text(s, LEFT, 3.45, CW, 1.2, [
    ("Round-2 proposal (runs on your sign-off):", 12, True),
    ("d ∈ {50, 200, 1,000, 5,000} · MLP mechanisms · gated edge subset\n"
     "Regime-GRACE vs blind PCMCI+ vs linear-additive u vs CD-NOD-style\n"
     "grouping above the ceiling, reported as its own curve", 12, False)],
    leading=1.5)
notes(s, "Her first 8/15 question. Be precise: the ~100 ceiling is GRACE's and "
         "GRACE only ran P2's slot graph; ridge ran E0/MINJA and is bounded by "
         "samples per regime, not variables. The mandatory scaling sim does not "
         "exist yet — decision 3.")

# ---------------- S5: E0 ----------------
s = add_slide("Pooling hides gated edges", 20, tag="Part 1 · E0 simulation")
pic(s, F_REGIME, 0.35, 0.72, 9.3)
add_text(s, LEFT, 3.62, CW, 1.3, [
    ("Blind discovery misses both read edges at 3,000 steps. With 20× data it "
     "finds them — but still cannot say when they act. Regime conditioning finds "
     "and gates both at 3,000.", 12, False),
    ("Detection is a sample-size issue; the gate is a model-class issue.",
     12.5, True)], leading=1.4)
notes(s, "If she asks power vs identifiability: the sweep answers it — blind gets "
         "detection at 20x data (0/2 -> 2/2), never the gate. Additive-u follows "
         "regime frequency: 0/2 on E0 (5% active) vs 2/3 on MINJA (69%). Our "
         "additive arms are linear; nonparametric CD-NOD is in the round-2 sim.")

# ---------------- S6: P2 ----------------
s = add_slide("Same graph, forward: compact memory (PASS)", 20,
              tag="Part 2 · MemoryArena, held-out IDs")
pic(s, F_P2, 0.55, 0.72, 9.0)
add_text(s, LEFT, 4.30, CW, 0.9, [
    ("Frozen criterion: PS loss ≤ 5 points and input reduction ≥ 30% vs noG → "
     "measured −1.43 and −39.55% → PASS.  vs BM25: +19.8 PS at −70% input.",
     12, False),
    ("Boundary: graph selection and compact serialization are bundled.",
     12, True)], leading=1.4)
notes(s, "Bars pool all persons; the frozen criterion uses episode-mean deltas "
         "(−1.43 vs pooled −1.52 — episodes differ in person count). Unbundling "
         "selection from serialization is decision 4.")

# ---------------- S7: P3 ----------------
s = add_slide("Same graph, backward: auditing works, online defense does not", 18,
              tag="Part 3 · safety")
pic(s, F_P3, 0.35, 0.72, 9.3)
add_text(s, LEFT, 3.68, CW, 1.3, [
    ("Auditing: the hidden read edge recovered in 3/3 seeds and traced to its "
     "write round (MINJA); 2/2 poison records recovered label-free (AgentPoison).",
     12, False),
    ("Online deletion gate: ASR drops on both carriers, but only 1/3 seed blocks "
     "improve → FAIL, kept.", 12, False),
    ("One-line caveats: MINJA's event channel is oracle-tagged; AgentPoison "
     "recovery had no pre-registered PASS.", 11, False, GRAY)], leading=1.4)
notes(s, "Keep the halves separate: auditing (her 8/21 trustworthiness pitch) is "
         "supported; mitigation failed its frozen rule twice. The 6/36 here is "
         "the mitigation round's own ungated arm; the replication round measured "
         "8/36 (appendix). Power audit says the AgentPoison FAIL is mostly the "
         "acceptance rule — decision 5.")

# ---------------- S8: decisions ----------------
s = add_slide("What we need from you", 20, tag="decisions")
add_text(s, LEFT, 0.85, CW, 3.9, [
    ("1 · Formulation — build the paper on write → hold → regime-gated read?",
     13, True),
    ("", 5, False),
    ("2 · Samples — does per-edge × lag × regime counting settle the sample "
     "question?", 13, True),
    ("", 5, False),
    ("3 · Scale — run the d = 50 → 5,000 simulation as round 2?", 13, True),
    ("", 5, False),
    ("4 · Boundaries — how much unbundling before submission? "
     "(P2 bundle · MINJA oracle tag)", 13, True),
    ("", 5, False),
    ("5 · P3-B — keep the FAILs and stop, or run the drafted round 3? "
     "(2/3 blocks had zero attacks to prevent)", 13, True),
    ("", 5, False),
    ("6 · Venue — which deadline? ICLR 2027: abstract Sep 18 / paper Sep 25.",
     13, True)], leading=1.35)
add_text(s, LEFT, 4.80, CW, 0.4, [
    ("Only 3 and 5 need new runs.", 12, False, GRAY)], leading=1.3)
notes(s, "1-2 close the formulation. 3 and 5 are the only new-compute items. 4 "
         "sets the writing bar. If she confirms ICLR 2027, Sep 18/25 makes 3 and "
         "5 this week's decisions. Round-3 draft is written and waits for "
         "sign-off; round-2 negatives stay pass=false regardless.")

# =====================================================================
# APPENDIX (dense backup)
# =====================================================================

s = add_slide("Appendix — P2: from PS = 0% to a frozen PASS", 17,
              tag="appendix · P2 audit trail")
add_text(s, LEFT, 0.80, CW, 4.2, [
    ("IDs 1–5 (original protocol):  PS 0.00% on both graph arms — reproduced "
     "exactly in the pre-registered rerun. Cause: the strict full-plan denominator "
     "(SPS 96.9% while one wrong slot fails the whole person); scorer and parsing "
     "verified clean.", 10.8, False),
    ("Failure audit:  31/37 persons failed only on slots the query never asked to "
     "change → deterministic query-target/base-inheritance decoder (no gold "
     "access, no evaluator in the loop).", 10.8, False),
    ("Same IDs, exploratory:  learned arm 100% PS online; learned−pure +18.69 "
     "points [+12.86, +28.33]. Labeled exploratory repair — developed after seeing "
     "the failures, so it stayed outside the confirmatory record.", 10.8, False),
    ("IDs 101–110 (held-out):  main judgement PASS — pure−BM25 +16.55 [+9.17, "
     "+25.24], W/T/L 8/2/0. The graph/noG secondary condition FAILED: input "
     "−14.7% against the frozen 30%.", 10.8, False),
    ("Diagnosis:  context compression outpaced API-input compression (shared "
     "system/tool schema and ReAct history dominate) → compact target-delta "
     "serialization v3, frozen on dev IDs 101–103 only.", 10.8, False),
    ("IDs 111–120 (new frozen round):  −1.43 PS / −39.55% input → PASS.",
     10.8, True)], leading=1.42)
footer2(s, "All intermediate negatives preserved unchanged in results/ · "
           "e2e_scores*.json · e2e_failure_audit*.json · "
           "e2e_scores_holdout_101_110.json · p2_compact_v3/round_summary.json")
notes(s, "The audit trail if she asks how a 0% became a PASS: reproduce the zero, "
         "audit the failures, freeze a decoder fix, fail the input criterion "
         "once, diagnose, freeze v3, pass on fresh held-out IDs.")

s = add_slide("Appendix — MINJA per-seed detail and boundaries", 17,
              tag="appendix · P3-A")
add_text(s, LEFT, 0.80, 4.55, 3.9, [
    ("Replication (3 seeds, 243 rounds):", 11.5, True),
    ("seed 0 ·  held-out ASR 6/12 · decisive cell 12/37 = 0.324 · edge found, "
     "gated", 10.5, False),
    ("seed 1 ·  ASR 1/12 · 6/31 = 0.194 · found, gated", 10.5, False),
    ("seed 2 ·  ASR 1/12 · 4/8 = 0.500 · found, without gate flag", 10.5, False),
    ("aggregate ·  ASR 8/36 = 0.222, Wilson [0.117, 0.381] · decisive 22/76 = "
     "0.289 [0.200, 0.400] · note-free control 0/77", 10.5, False),
    ("(the 6/36 on the Part 3 slide is the later P3-B mitigation round's own "
     "ungated arm — a separate run from this 8/36)", 9.5, False, GRAY),
    ("Discovery baselines (same traces):", 11.5, True),
    ("blind PCMCI 0/3 · additive u-as-node 2/3 · regime-subsampled 1/3 · seed-0 "
     "read coefficient 0.77 in-regime (p = 2.9e-6) vs 0.02 outside (p = 0.85)",
     10.5, False)], leading=1.34)
add_text(s, 5.15, 0.80, 4.5, 3.9, [
    ("Boundaries on the claim:", 11.5, True),
    ("·  poison_retr is built from the record's is_poison tag: the recovered "
     "structure sits on an oracle-tagged event channel, and MINJA record "
     "discovery stays outside the claim.", 10.5, False),
    ("·  lag-0 orientation comes from the instrumented retrieve-before-act order.",
     10.5, False),
    ("·  seed heterogeneity is real (test ASR 6/12 vs 1/12): the pathway "
     "reproduces, its strength is seed-dependent.", 10.5, False),
    ("·  MemAudit reproduction scores CMIS macro AUC 0.866 / precision@k 0.874; "
     "our claim is temporal ancestry plus one-graph-two-uses.", 10.5, False)],
    leading=1.34)
footer2(s, "minja_replication_summary.json · minja_audit_report_{v2,seed1,seed2}"
           ".json · memaudit_baseline_v2.json (lexical overlap stands in for NLI)")
notes(s, "Per-seed backup. The additive-u 2/3 here vs 0/2 on E0 is regime "
         "frequency (trigger active 69% of MINJA rounds vs 5% on E0). The "
         "boundaries column keeps P3-A at Layer 1.")

s = add_slide("Appendix — the regime condition in the replay gate", 17,
              tag="appendix · offline coverage")
pic(s, F_GATE, 0.50, 0.75, 9.0)
add_text(s, LEFT, 3.15, CW, 1.4, [
    ("Dropping the regime condition keeps coverage (+2 anomalous rounds) while "
     "touching 43 more benign rounds — collateral 8.8% → 40.4%. The score-based "
     "baseline with the same regime condition is precise but covers only 36.4%.",
     11, False),
    ("Coverage counts on saved trajectories, zero LLM re-invocation — a "
     "policy-design signal. The online test of the same idea is P3-B, which "
     "failed.", 11, False)], leading=1.3)
footer2(s, "results/real/causal_gate_v2.json · three seeds, micro-pooled: 107 "
           "anomalous / 136 benign rounds")
notes(s, "The gate's operational value without overclaiming: replay-only "
         "coverage, clearly separated from the failed online test.")

s = add_slide("Appendix — why P3-B failed; the drafted round 3", 17,
              tag="appendix · P3-B")
add_text(s, LEFT, 0.80, 4.55, 3.9, [
    ("Power audit (diagnostic only):", 11.5, True),
    ("·  AgentPoison: seed blocks 0 and 2 had ungated 0/24 — zero attacks to "
     "prevent, so “≥2/3 blocks improve” was unsatisfiable with one evaluable "
     "block. Conditional on evaluability the gate went 1/1, with 3 touched "
     "preventions and 0 reverse triggers: this FAIL is mostly the acceptance "
     "rule.", 10.5, False),
    ("·  MINJA: a genuine limitation. ~55% of memory records never surfaced in "
     "calibration retrieval (poison-recall cap 0.24–0.39), and untouched controls "
     "flip answers at the size of the measured effect — clean-gated changed 10/24 "
     "answers while deleting zero retrievals.", 10.5, False)], leading=1.34)
add_text(s, 5.15, 0.80, 4.5, 3.9, [
    ("Round-3 draft (freezes only after sign-off):", 11.5, True),
    ("·  a-priori evaluability: a block counts only with ungated attacks ≥ 3; "
     "≥ 3 evaluable blocks by construction, predicted from earlier rounds' "
     "ungated ASR only.", 10.5, False),
    ("·  exposure: longer calibration (≥ 40 eligible trigger rounds) + the frozen "
     "embedding-cluster expansion already validated on AgentPoison (recall 0.50 → "
     "1.00).", 10.5, False),
    ("·  noise: primary endpoint becomes touched-only net prevention against a "
     "pre-registered flip-rate floor from round 2; block ASR turns secondary.",
     10.5, False),
    ("·  round-2 verdicts stay pass=false whatever round 3 finds.", 10.5, True)],
    leading=1.34)
footer2(s, "results/real/p3b_power_audit.json · "
           "docs/p3b-round3-preregistration-draft-2026-08-30.md · "
           "results/diagnostics/e0_blind_power_sweep.json")
notes(s, "Backup for decision 5. The carriers failed for different reasons: "
         "AgentPoison's rule was unsatisfiable; MINJA's limits are exposure and "
         "call noise, both measured. The draft fixes each with a frozen rule.")

out = os.path.join(HERE, "8-30-yujia.pptx")
prs.save(out)
print("saved", out, "with", len(prs.slides._sldIdLst), "slides")
