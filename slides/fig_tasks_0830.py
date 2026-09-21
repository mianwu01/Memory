#!/usr/bin/env python3
"""Draw the two task figures and place them on the 'pieces' slides of 8-30.pptx.

fig_task_p2: a MemoryArena travel query — the itinerary grid, the queried cell,
the graph ancestor that gets serialized with it, everything else inherited.
fig_task_p3: the MINJA attack as write -> hold -> gated read, with the backward
audit arc that P3-A must recover.

One-shot patch: finds the slides titled 'Request 1 — the pieces' and
'Request 2 — the pieces', removes their body text, inserts figure + condensed
text. Safe to re-run (it rebuilds those two slides' bodies each time).
"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figs_0830")
BLACK = RGBColor(0, 0, 0)
BODY_FONT = "Georgia"
BLUE = "#2a78d6"
LIGHTBLUE = (0.88, 0.92, 0.99)

plt.rcParams.update({
    "figure.facecolor": "white", "axes.facecolor": "white",
    "text.color": "black", "font.size": 10,
    "font.family": "serif", "font.serif": ["Georgia", "DejaVu Serif"],
})


def fig_task_p2():
    f, ax = plt.subplots(figsize=(9.2, 2.7))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")

    def box(x, y, w, h, txt, fc="white", ec="black", lw=1.2, fs=8.6):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.008",
                                    fc=fc, ec=ec, lw=lw))
        if txt:
            ax.text(x + w / 2, y + h / 2, txt, ha="center", va="center",
                    fontsize=fs)

    def arrow(x1, y1, x2, y2, lw=1.3, color="black", ls="-", rad=0.0):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
                                     mutation_scale=12, lw=lw, color=color,
                                     ls=ls, connectionstyle=f"arc3,rad={rad}"))

    # query bubble
    box(0.01, 0.52, 0.165, 0.30,
        "round-t query:\n“move Alice's\nday-2 dinner to X”", fc=LIGHTBLUE, fs=8.8)

    # itinerary grid: one person, 2 days x 7 slots
    slots = ["city", "transp.", "bkfst", "attr.", "lunch", "dinner", "hotel"]
    x0, cw, chh = 0.255, 0.068, 0.155
    ax.text(x0 + 3.5 * cw, 0.985, "Alice's itinerary (persons × days × 7 slots)",
            ha="center", fontsize=8.6, style="italic", color="0.3")
    for c, sname in enumerate(slots):
        ax.text(x0 + c * cw + cw / 2, 0.90, sname, ha="center", fontsize=7.4,
                color="0.25")
    for r, (dy, dlab) in enumerate([(0.68, "day 1"), (0.42, "day 2")]):
        ax.text(x0 + 7 * cw + 0.006, dy + chh / 2, dlab, ha="left", va="center",
                fontsize=8.2)
        for c in range(7):
            fc, lw = "white", 1.0
            if r == 1 and c == 5:              # queried target cell
                fc, lw = BLUE, 1.4
            elif r == 0 and c == 5:            # graph ancestor
                fc, lw = LIGHTBLUE, 1.4
            box(x0 + c * cw, dy, cw - 0.006, chh, "", fc=fc, lw=lw)
    arrow(0.178, 0.62, 0.248, 0.50, rad=-0.15)
    # learned edge day1-dinner -> day2-dinner, label in the inter-row gap
    dx = x0 + 5 * cw + (cw - 0.006) / 2
    arrow(dx, 0.675, dx, 0.582, lw=1.6, color=BLUE)
    ax.text(dx - 0.014, 0.628, "learned edge (lag 1)", fontsize=7.4,
            color=BLUE, va="center", ha="right")

    # legend
    lx = 0.80
    for dy, fc, lab in [(0.80, BLUE, "query target — sent to the LLM"),
                        (0.68, LIGHTBLUE, "graph ancestor — sent too"),
                        (0.56, "white", "everything else — inherited\nfrom the public base, never sent")]:
        box(lx, dy, 0.030, 0.075, "", fc=fc)
        ax.text(lx + 0.042, dy + 0.037, lab, fontsize=7.8, va="center")

    # bottom flow
    box(0.10, 0.05, 0.22, 0.20, "LLM plans only the\ntarget cells", fs=8.4)
    arrow(0.325, 0.15, 0.385, 0.15)
    box(0.39, 0.05, 0.24, 0.20, "decoder inherits the rest\nfrom the base plan", fs=8.4)
    arrow(0.635, 0.15, 0.695, 0.15)
    box(0.70, 0.05, 0.26, 0.20,
        "official scorer:\nPS (whole person) · SPS (slot) · SR", fs=8.4)
    f.tight_layout(pad=0.2)
    p = os.path.join(FIG, "task_p2.png")
    f.savefig(p, dpi=200, bbox_inches="tight"); plt.close(f); return p


def fig_task_p3():
    f, ax = plt.subplots(figsize=(9.2, 2.1))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")

    def box(x, y, w, h, txt, fc="white", lw=1.3, fs=8.8):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.010",
                                    fc=fc, ec="black", lw=lw))
        ax.text(x + w / 2, y + h / 2, txt, ha="center", va="center", fontsize=fs)

    def arrow(x1, y1, x2, y2, lw=1.4, color="black", ls="-", rad=0.0):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
                                     mutation_scale=13, lw=lw, color=color,
                                     ls=ls, connectionstyle=f"arc3,rad={rad}"))

    y, h = 0.50, 0.32
    box(0.01, y, 0.195, h, "attacker's normal-looking\nquery + shortcut note")
    arrow(0.208, y + h / 2, 0.252, y + h / 2)
    ax.text(0.23, y + h + 0.06, "write", ha="center", fontsize=8,
            style="italic")
    box(0.255, y, 0.175, h, "poison record\nin agent memory", fc=LIGHTBLUE)
    arrow(0.433, y + h / 2, 0.545, y + h / 2)
    ax.text(0.49, y + h + 0.06, "hold — benign rounds, record dormant",
            ha="center", fontsize=8, style="italic")
    box(0.548, y, 0.155, h, "victim query\nwith trigger")
    arrow(0.706, y + h / 2, 0.788, y + h / 2, lw=2.0, color=BLUE)
    ax.text(0.747, y + h + 0.06, "read (retrieval)", ha="center",
            fontsize=8, style="italic", color=BLUE)
    box(0.792, y, 0.185, h, "anomalous answer", fc=LIGHTBLUE)

    # backward audit arc, bowed below the boxes
    arrow(0.86, y - 0.04, 0.36, y - 0.04, lw=1.4, color=BLUE,
          ls=(0, (4, 2)), rad=-0.22)
    ax.text(0.615, 0.045, "P3-A audit: trace the answer back through retrieval "
            "to the write round", ha="center", fontsize=8.6, color=BLUE)
    ax.text(0.015, 0.10, "= the object,\nadversarially\ninstantiated",
            fontsize=8.2, style="italic", color="0.3", va="center")
    f.tight_layout(pad=0.2)
    p = os.path.join(FIG, "task_p3.png")
    f.savefig(p, dpi=200, bbox_inches="tight"); plt.close(f); return p


F_P2 = fig_task_p2()
F_P3 = fig_task_p3()

# ---------------- patch the deck ----------------
SRC = os.path.join(HERE, "8-30.pptx")
prs = Presentation(SRC)


def add_text(slide, y, lines, leading=1.35):
    tb = slide.shapes.add_textbox(Inches(0.38), Inches(y), Inches(9.24), Inches(1.6))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (txt, sz, bold) in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.line_spacing = leading
        r = para.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold
        r.font.color.rgb = BLACK; r.font.name = BODY_FONT
    return tb


def rebuild(title_key, fig, fig_y, fig_w, text_y, lines):
    for s in prs.slides:
        titles = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        if any(title_key in t for t in titles):
            for sh in list(s.shapes):
                if sh.has_text_frame and sh.top / 914400 > 0.45:
                    sh._element.getparent().remove(sh._element)
                elif sh.shape_type == 13:
                    sh._element.getparent().remove(sh._element)
            s.shapes.add_picture(fig, Inches((10 - fig_w) / 2), Inches(fig_y),
                                 width=Inches(fig_w))
            add_text(s, text_y, lines)
            return True
    return False


ok1 = rebuild(
    "Request 1 — the pieces", F_P2, 0.72, 9.2, 3.70,
    [("GRACE (causalts): neural temporal discovery, ~100-variable capacity, "
      "wrapped multi-trial — episodes are independent trials, no cross-episode "
      "lags.", 11, False),
     ("Our module: the learned 7-slot graph picks what to send; a "
      "runtime-registered HTTP memory API plugs it into the arena beside 13 "
      "built-in memory systems.", 11, False)])

ok2 = rebuild(
    "Request 2 — the pieces", F_P3, 0.80, 9.2, 3.05,
    [("MINJA harness: 81 rounds per seed, instrumented memory, 6 event channels "
      "per round (trigger, note, poison-in-memory, poison-retrieved, anomalous, "
      "correct).", 11, False),
     ("AgentPoison-StrategyQA: an optimized trigger pulls 2 poisoned records "
      "from a frozen DPR index (9,253 passages) → “I don't know” failures; "
      "paired ungated / no-op / gated arms from one immutable snapshot.",
      11, False),
     ("Gates: the recovered ancestry + regime decide which records to delete — "
      "offline replay first, then a real online intervention.", 11, False)])

assert ok1 and ok2, (ok1, ok2)
prs.save(SRC)
print("patched", SRC)
