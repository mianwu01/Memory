#!/usr/bin/env python3
"""One-shot: insert the two-request narrative slides into slides/8-30.pptx.

Keeps every existing (hand-edited) slide untouched. Adds, per request, three
slides in the order of the agreed explanation — how we picked the setup, what
the pieces are, what we did — and places them directly before that request's
existing results slide. Original file is backed up to 8-30.pptx.bak once.

Expected starting deck (8 slides): 0 title · 1 object · 2 SCM · 3 scale ·
4 E0 · 5 P2 results · 6 P3 results · 7 appendix replay gate.
Final order: 0 1 2 3 4 · R1-choice R1-pieces R1-did · 5 · R2-choice R2-pieces
R2-did · 6 · 7   (14 slides).
"""
import os
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "8-30.pptx")
BAK = SRC + ".bak"
if not os.path.exists(BAK):
    shutil.copy2(SRC, BAK)

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x55, 0x55, 0x55)
BODY_FONT = "Georgia"

prs = Presentation(SRC)
assert len(prs.slides._sldIdLst) == 8, "expected the 8-slide hand-edited deck"
BLANK = ([l for l in prs.slide_layouts if l.name.upper() == "BLANK"] or
         [prs.slide_layouts[6]])[0]
LEFT, CW = 0.38, 9.24


def add_slide(title, title_size=20, tag=None):
    s = prs.slides.add_slide(BLANK)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    tb = s.shapes.add_textbox(Inches(LEFT), Inches(0.18), Inches(CW), Inches(0.55))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = title; r.font.size = Pt(title_size); r.font.bold = True
    r.font.color.rgb = BLACK; r.font.name = BODY_FONT
    if tag:
        tb = s.shapes.add_textbox(Inches(6.4), Inches(0.02), Inches(3.35), Inches(0.3))
        p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = tag
        r.font.size = Pt(10); r.font.italic = True
        r.font.color.rgb = GRAY; r.font.name = BODY_FONT
    return s


def add_text(slide, y, lines, leading=1.4):
    tb = slide.shapes.add_textbox(Inches(LEFT), Inches(y), Inches(CW), Inches(4.4))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.line_spacing = leading
        if isinstance(line, tuple) and len(line) == 4:
            txt, sz, bold, col = line
        else:
            txt, sz, bold = line; col = BLACK
        r = para.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold
        r.font.color.rgb = col; r.font.name = BODY_FONT
    return tb


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------- Request 1 ----------------
s = add_slide("Request 1 — choosing the benchmark", tag="effectiveness · why MemoryArena")
add_text(s, 0.80, [
    ("Your constraints: an independent module injected into an existing repo "
     "(7/29) · one concrete benchmark, variables and samples defined first (8/15).",
     11.5, False),
    ("", 5, False),
    ("·  LongMemEval — no room for a graph on recall (ancestors = the answers; "
     "BM25 already 0.947) → diagnostic only", 11, False),
    ("·  MemoryAgentBench — setup cost misjudged → dropped", 11, False),
    ("·  MemoryArena — persistent state drives later actions in all 4 "
     "environments; plug-in memory registry, official scorer, 13 built-in "
     "systems → chosen", 11, True),
    ("", 5, False),
    ("Travel environment: its multi-round correction loop is the closest real "
     "write → hold → read. Upstream has no license, so our system registers at "
     "runtime and the upstream stays clean.", 11.5, False),
    ("Discovery method, verified by implementation: GRACE (causalts) main · "
     "CDNOTS+ baseline · UnCLe / CUTS+ / AVICI excluded (d ≤ 8 / no lag / no "
     "time series).", 11.5, False)])
notes(s, "The choice story: her constraints first, then the three candidates and "
         "why MemoryArena won — real persistent-state structure plus a plug-in "
         "registry so we compare against its built-in memory systems on the "
         "official scorer.")
R1_CHOICE = s

s = add_slide("Request 1 — the pieces", tag="components")
add_text(s, 0.85, [
    ("MemoryArena travel:  one episode = a travel group; each person = a "
     "multi-day itinerary × 7 slots (city, transport, breakfast, attraction, "
     "lunch, dinner, hotel). Queries ask for specific changes. Official scores: "
     "PS = whole person correct · SPS = slot level · SR = episode.", 11.5, False),
    ("", 6, False),
    ("GRACE (causalts):  neural temporal causal discovery, ~100-variable "
     "capacity. We wrap it multi-trial: episodes are independent trials, and no "
     "lag crosses an episode boundary.", 11.5, False),
    ("", 6, False),
    ("Our module:  a learned 7-slot graph; the memory backend serializes only "
     "the query's cells plus their graph ancestors · registered at runtime over "
     "an HTTP memory API · a deterministic decoder inherits unqueried slots from "
     "the public base plan.", 11.5, False)])
notes(s, "Plain description of the three components: the benchmark's structure "
         "and scores, what GRACE is and the multi-trial wrapper, and what our "
         "module actually serializes.")
R1_PIECES = s

s = add_slide("Request 1 — what we did", tag="three stages, all kept")
add_text(s, 0.85, [
    ("1 ·  IDs 1–5: PS 0% on both graph arms while SPS ≈ 97% — reproduced in a "
     "pre-registered rerun. The audit showed 31/37 persons failed only on slots "
     "the query never asked to change → built the inheritance decoder "
     "(exploratory repair on the same IDs: 100% PS).", 11.5, False),
    ("", 6, False),
    ("2 ·  Held-out IDs 101–110: beat BM25 by +16.55 PS (main judgement PASS), "
     "but graph-vs-noG input fell only 14.7% against the required 30% → "
     "secondary FAIL. Diagnosis: shared scaffold and history dominate API input "
     "→ compact target-delta serialization.", 11.5, False),
    ("", 6, False),
    ("3 ·  New frozen held-out IDs 111–120 — the result on the next slide.",
     11.5, True)])
notes(s, "The journey in three stages, negatives preserved: reproduce the zero, "
         "audit and fix the decoder, fail the input criterion once, freeze the "
         "compact serialization, then the new frozen round.")
R1_DID = s

# ---------------- Request 2 ----------------
s = add_slide("Request 2 — choosing the attack", tag="trustworthiness · why MINJA + AgentPoison")
add_text(s, 0.78, [
    ("Your two routes (8/21): modify an existing attack and recover something "
     "hidden, or construct scenarios ourselves. We did both — constructed = E0; "
     "existing = a real memory-poisoning attack, so the hidden driver has ground "
     "truth.", 11.5, False),
    ("", 5, False),
    ("·  MINJA — chosen first: self-contained QA memory injection, reproducible "
     "on our API model", 11, True),
    ("·  AgentPoison — second carrier via its ReAct-StrategyQA track (enables "
     "label-free recovery); the EHR-agent track is off-limits", 11, False),
    ("·  MemAudit — reimplemented as the baseline; our claim is temporal "
     "ancestry + one-graph-two-uses, a detection-score win is out of scope",
     11, False),
    ("", 5, False),
    ("A detour that shaped the method: with diluted benign filler the driver "
     "vanished (0 note-free anomalous rounds) — retrieval was confounded with "
     "the visible note. The faithful config reproduced it, and recovery is only "
     "declared on the note-free ∧ poison-retrieved cell.", 11.5, False)])
notes(s, "Why these carriers: MINJA is reproducible and self-contained; "
         "AgentPoison's StrategyQA track adds the label-free test; the EHR track "
         "is off-limits. The dilution detour is worth telling — it is why the "
         "decisive-cell rule exists.")
R2_CHOICE = s

s = add_slide("Request 2 — the pieces", tag="components")
add_text(s, 0.85, [
    ("MINJA:  the attacker poisons an agent's memory through normal queries; "
     "later, trigger-bearing queries retrieve those records and shift answers. "
     "Our harness runs 81 rounds per seed with instrumented memory, logging 6 "
     "event channels per round (trigger, note, poison-in-memory, "
     "poison-retrieved, anomalous, correct).", 11.5, False),
    ("", 6, False),
    ("AgentPoison-StrategyQA:  an optimized trigger pulls 2 poisoned records "
     "from a frozen DPR index (9,253 passages) and induces “I don't know” "
     "failures. Paired ungated / no-op / gated arms run from one immutable "
     "snapshot.", 11.5, False),
    ("", 6, False),
    ("Gates:  the recovered ancestry + regime decide which records to delete — "
     "replayed offline first, then a real online intervention.", 11.5, False)])
notes(s, "What each attack actually is, in one breath each, plus the gate that "
         "turns recovery into a defense attempt.")
R2_PIECES = s

s = add_slide("Request 2 — what we did", tag="replicate → recover → defend")
add_text(s, 0.85, [
    ("Replicate:  3 MINJA seeds × 81 rounds. The pathway is real — decisive "
     "note-free ∧ poison-retrieved cell 22/76 = 0.289, note-free control 0/77, "
     "held-out ASR 8/36 — with large seed differences (6/12 vs 1/12).",
     11.5, False),
    ("", 6, False),
    ("Recover:  Regime-GRACE and the discovery baselines on the same traces; on "
     "AgentPoison, a label-free driver (64 calibration trajectories, 288 "
     "held-out; poison labels attached only afterwards).", 11.5, False),
    ("", 6, False),
    ("Defend:  replayed the ancestry + regime gate offline, then froze two "
     "online intervention protocols — results on the next slide.", 11.5, True)])
notes(s, "Three verbs: replicate the attack, recover the driver, attempt the "
         "defense. The next slide carries both verdicts.")
R2_DID = s

# ---------------- reorder ----------------
lst = prs.slides._sldIdLst
ids = list(lst)                      # 0..7 original, 8..13 new
order = [ids[0], ids[1], ids[2], ids[3], ids[4],
         ids[8], ids[9], ids[10],    # R1 choice / pieces / did
         ids[5],                     # P2 results
         ids[11], ids[12], ids[13],  # R2 choice / pieces / did
         ids[6],                     # P3 results
         ids[7]]                     # appendix
for el in order:
    lst.remove(el)
for el in order:
    lst.append(el)

prs.save(SRC)
print("saved", SRC, "with", len(prs.slides._sldIdLst), "slides (backup: 8-30.pptx.bak)")
